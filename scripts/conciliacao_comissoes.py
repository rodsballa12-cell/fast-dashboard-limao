# -*- coding: utf-8 -*-
"""Confere comissão devida (Trinks) × comissão paga (Pix no razão da conta XP), por profissional e quinzena.

Por que existe: em 16/10/2026 o Rodrigo pediu para conferir o que o financeiro pagou
contra o que deveria pagar, "principalmente as comissões", e automatizar a conferência.
No total a equipe bateu com o Trinks (R$ 38.249 pagos × R$ 38.089 devidos, 23/07–30/09),
mas por pessoa havia desvios — a maior parte por cadastro (profissional desligada
EXCLUÍDA do Trinks, produção sem nome) e alguns a explicar.

Como calcula:
- Devido = cada serviço do Trinks (preço; serviço de pacote pelo preço de tabela) ×
  taxa da categoria em data/comissoes.json (Spa: 25% uniforme). Mesma regra do ranking
  de comissões do painel (github_refresh.py).
- Pago = Pix do razão (aba Conta_XP do Excel do PC) com categoria "OPEX (Comissão)".
- Janela de pagamento da quinzena: do 1º dia da quinzena até 5 dias depois do fim
  (os repasses saem no dia 15 e no último dia, às vezes no 1º dia útil seguinte).
- Nome do Pix × nome do Trinks: casa por 1º + último nome; apelidos e cadastros
  duplicados (ex.: MEI) em comissoes_aliases.json, ao lado do Excel. Profissionais
  desligadas usam o nome de data/prof_overrides.json quando preenchido.

PRIVACIDADE: o repositório é público. O resultado (valor pago por pessoa) NUNCA vai para
o repositório — sai só no arquivo de pendências do OneDrive (aba "11 Comissões x Trinks")
e num JSON ao lado dele. Este script não grava nada em data/.

Uso:
    python scripts/conciliacao_comissoes.py            # grava a aba se o resultado mudou
    python scripts/conciliacao_comissoes.py --forcar   # regrava mesmo sem mudança
    python scripts/conciliacao_comissoes.py --seco     # só imprime, não grava

Saída em ASCII de propósito (log da tarefa agendada passa pela codepage do console).
"""
from __future__ import annotations

import collections
import datetime as dt
import hashlib
import json
import re
import shutil
import sys
import tempfile
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAIZ = Path(__file__).resolve().parents[1]
PASTA = Path(r"C:\Users\rods_\OneDrive\Franquia - FAST\Claude")
PAINEL = PASTA / "Painel_Gestao_Financeira_SIIBELLO.xlsx"
PEND = PASTA / "Pendencias_Identificacao_SIIBELLO.xlsx"
ALIASES = PASTA / "comissoes_aliases.json"
RESULTADO = PASTA / "conciliacao_comissoes.json"
ABA = "11 Comissões x Trinks"
INICIO = dt.date(2026, 7, 23)          # 1ª venda da Escova no Trinks
LIMIAR_RS, LIMIAR_PCT = 100.0, 0.05     # diferença que vira pendência
FOLGA_PAGTO = 5                         # dias após o fim da quinzena


def ascii_(s: str) -> str:
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()


def norm(s) -> str:
    s = ascii_(s).lower()
    s = re.sub(r"\(.*?\)", "", s)
    s = re.sub(r"^\d{2}\.\d{3}\.\d{3}\s+", "", s.strip())
    s = re.sub(r"\d{11}", "", s)
    return re.sub(r"\s+", " ", s).strip()


def quinzena(d: dt.date) -> tuple[int, int, int]:
    return (d.year, d.month, 1 if d.day <= 15 else 2)


def limites(q: tuple[int, int, int]) -> tuple[dt.date, dt.date]:
    y, m, h = q
    if h == 1:
        return dt.date(y, m, 1), dt.date(y, m, 15)
    prox = dt.date(y + (m == 12), m % 12 + 1, 1)
    return dt.date(y, m, 16), prox - dt.timedelta(days=1)


def rotulo(q) -> str:
    return f"{'1a' if q[2] == 1 else '2a'} {q[1]:02d}/{q[0]}"


def ler_json(p: Path, padrao=None):
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return padrao


# ---------------------------------------------------------------- devido (Trinks)
def devido_trinks():
    regras = ler_json(RAIZ / "data" / "comissoes.json", {})
    taxa = {k.lower().strip(): v for k, v in (regras.get("por_categoria") or {}).items()}
    padrao = regras.get("regra_padrao_pct", 0.30)
    over = (ler_json(RAIZ / "data" / "prof_overrides.json", {}) or {}).get("profs", {}) or {}
    dev = collections.defaultdict(lambda: collections.defaultdict(float))
    nomes: dict[str, str] = {}
    ate = INICIO
    for pasta, spa in ((RAIZ / "data", False), (RAIZ / "data" / "spa", True)):
        pm = ((ler_json(pasta / "profissionais_cache.json", {}) or {}).get("payload") or {}).get("prof_map", {})
        cat_preco = {s.get("id"): float(s.get("preco") or 0)
                     for s in ((ler_json(pasta / "servicos_cache.json", {}) or {}).get("payload") or [])}
        for tx in ((ler_json(pasta / "transacoes_ano_cache.json", {}) or {}).get("payload") or []):
            d = dt.datetime.fromisoformat(tx["dataReferencia"]).date()
            ate = max(ate, d)
            for s in tx.get("servicos") or []:
                pid = s.get("idProfissionalQueRealizouServico")
                if not pid:
                    continue
                nome = pm.get(str(pid)) or over.get(str(pid)) or f"Desligada #{pid}"
                v = float(s.get("preco") or 0) or cat_preco.get(s.get("id"), 0)
                t = 0.25 if spa else taxa.get((s.get("categoria") or "").lower().strip(), padrao)
                k = norm(nome)
                nomes[k] = nome
                dev[k][quinzena(d)] += v * t
    return dev, nomes, ate


# ---------------------------------------------------------------- pago (razão XP)
def pagos_razao():
    # Lê uma CÓPIA do Excel: não trava a planilha nem disputa com o Excel aberto.
    tmp = Path(tempfile.gettempdir()) / "painel_conc_comissoes.xlsx"
    shutil.copy2(PAINEL, tmp)
    x = openpyxl.load_workbook(tmp, data_only=True, read_only=True)["Conta_XP"]
    pagos = collections.defaultdict(list)
    for row in x.iter_rows(min_row=16, max_col=7, values_only=True):
        a, b, _c, v, _e, _f, g = row
        if isinstance(a, dt.datetime) and isinstance(v, (int, float)) and v < 0 and str(g).startswith("OPEX (Comiss"):
            pagos[norm(b)].append((a.date(), round(-v, 2)))
    return pagos


def casar(nome_pix: str, devidos, aliases) -> str | None:
    if nome_pix in aliases:
        return norm(aliases[nome_pix])
    tp = nome_pix.split()
    for kt in devidos:
        tt = kt.split()
        if tp and tt and tp[0] == tt[0] and (tp[-1] == tt[-1] or (len(tp) > 1 and len(tt) > 1 and tp[1] == tt[1])):
            return kt
    return None


def conciliar():
    dev, nomes, ate = devido_trinks()
    al = ler_json(ALIASES, {}) or {}
    aliases = {norm(k): v for k, v in (al.get("pix_para_trinks") or {}).items()}
    sem_regra = {norm(k): v for k, v in (al.get("pagos_sem_trinks") or {}).items()}   # gerente etc.
    pagos = pagos_razao()
    por = collections.defaultdict(list)
    for kp, ps in pagos.items():
        alvo = casar(kp, dev, aliases) or kp
        if alvo not in nomes:
            nomes[alvo] = kp.title()
        por[alvo].extend(ps)
    # quinzenas fechadas: fim + folga já passou e o Trinks cobre o fim da quinzena
    hoje = dt.date.today()
    qs = sorted({q for d in dev.values() for q in d})
    fechadas = [q for q in qs if limites(q)[1] + dt.timedelta(days=FOLGA_PAGTO) <= hoje and limites(q)[1] <= ate]
    linhas = []
    for k in sorted(set(dev) | set(por)):
        porq = []
        for q in fechadas:
            ini, fim = limites(q)
            fim_p = fim + dt.timedelta(days=FOLGA_PAGTO)
            du = round(dev.get(k, {}).get(q, 0), 2)
            pa = round(sum(v for d, v in por.get(k, []) if ini <= d <= fim_p and not (d > fim and quinzena(d) != q and d.day > FOLGA_PAGTO)), 2)
            porq.append(dict(q=rotulo(q), devido=du, pago=pa))
        du_t = round(sum(dev.get(k, {}).get(q, 0) for q in fechadas), 2)
        fim_t = limites(fechadas[-1])[1] + dt.timedelta(days=FOLGA_PAGTO) if fechadas else INICIO
        pa_t = round(sum(v for d, v in por.get(k, []) if d <= fim_t), 2)
        if du_t < 1 and pa_t < 1:
            continue
        dif = round(pa_t - du_t, 2)
        if k in sem_regra:
            sit, perg = "Sem base no Trinks", sem_regra[k]
        elif nomes.get(k, "").startswith("Desligada #"):
            sit, perg = "Cadastro excluido no Trinks", "Quem e? Preencher o nome em data/prof_overrides.json"
        elif du_t < 1:
            sit, perg = "Pix sem producao no Trinks", "E profissional desligada? Qual o cadastro dela no Trinks? (ou apelido em comissoes_aliases.json)"
        elif pa_t < 1:
            sit, perg = "Producao sem Pix de comissao", "Ainda nao pago? Recebe de outra forma (CLT, outro nome)?"
        elif abs(dif) <= max(LIMIAR_RS, LIMIAR_PCT * du_t):
            sit, perg = "OK", ""
        else:
            sit, perg = ("Pago acima do Trinks" if dif > 0 else "Pago abaixo do Trinks"), "Explicar (garantia minima, vale, adiantamento, desconto, cancelamento?)"
        pix = [f"{d:%d/%m} {v:.2f}" for d, v in sorted(por.get(k, [])) if d <= fim_t]
        linhas.append(dict(nome=nomes.get(k, k), devido=du_t, pago=pa_t, dif=dif, situacao=sit, pergunta=perg, quinzenas=porq, pix=pix))
    linhas.sort(key=lambda r: (r["situacao"] == "OK", -abs(r["dif"])))
    equipe = [r for r in linhas if r["situacao"] != "Sem base no Trinks"]
    tot = dict(devido=round(sum(r["devido"] for r in equipe), 2), pago=round(sum(r["pago"] for r in equipe), 2),
               sem_base=round(sum(r["pago"] for r in linhas if r["situacao"] == "Sem base no Trinks"), 2))
    tot["dif"] = round(tot["pago"] - tot["devido"], 2)
    return dict(gerado_em=dt.datetime.now().isoformat(timespec="seconds"), trinks_ate=ate.isoformat(),
                periodo=f"{INICIO:%d/%m/%Y} a {limites(fechadas[-1])[1]:%d/%m/%Y}" if fechadas else "-",
                quinzenas=[rotulo(q) for q in fechadas], totais=tot, linhas=linhas)


# ---------------------------------------------------------------- aba no arquivo de pendências
def gravar_aba(res):
    wb = openpyxl.load_workbook(PEND)
    respostas = {}
    if ABA in wb.sheetnames:   # preserva o que o Lucas já respondeu (coluna "Resposta")
        old = wb[ABA]
        for row in old.iter_rows(min_row=1, values_only=True):
            if row and row[0] and len(row) >= 13 and row[12]:
                respostas[str(row[0])] = row[12]
        idx = wb.sheetnames.index(ABA)
        del wb[ABA]
        ws = wb.create_sheet(ABA, idx)
    else:
        ws = wb.create_sheet(ABA)
    ws.sheet_view.showGridLines = False
    qs = res["quinzenas"][-5:]
    larg = [40, 13, 13, 12] + [11] * len(qs) + [26, 46, 52, 30]
    for i, w in enumerate(larg, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    F = {k: PatternFill("solid", fgColor=v) for k, v in dict(hdr="1F3864", tot="D9E1F2", ok="E2EFDA", warn="FFF2CC", bad="FCE4D6", amar="FFFF99").items()}
    fino = Side(style="thin", color="BFBFBF"); B = Border(left=fino, right=fino, top=fino, bottom=fino)
    NUM = '#,##0.00;[Red]-#,##0.00;"–"'

    def c(r, col, v, bold=False, fill=None, num=False, color=None, size=None, borda=True):
        q = ws.cell(r, col, v)
        q.font = Font(bold=bold, color=color, size=size)
        if fill: q.fill = F[fill]
        if num: q.number_format = NUM
        q.alignment = Alignment(wrap_text=True, vertical="center")
        if borda: q.border = B
        return q

    ncol = len(larg)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c(1, 1, "COMISSÕES — devido pelo Trinks × pago por Pix (conferência automática)", bold=True, size=14, color="1F3864", borda=False)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    c(2, 1, f"Período {res['periodo']} (quinzenas fechadas) · Trinks até {res['trinks_ate']} · gerado em {res['gerado_em'][:16].replace('T', ' ')} "
            f"por scripts/conciliacao_comissoes.py (roda todo dia no agente). Devido = serviço × taxa da categoria (data/comissoes.json; Spa 25%). "
            f"Pago = Pix 'OPEX (Comissão)' do razão. Diferença acima de R$ {LIMIAR_RS:.0f} e de {LIMIAR_PCT:.0%} vira pendência.", color="595959", borda=False)
    ws.row_dimensions[2].height = 44
    t = res["totais"]; r = 4
    for rot, v, f in (("Comissão devida à equipe (Trinks)", t["devido"], None), ("Comissão paga à equipe (Pix)", t["pago"], None),
                      ("DIFERENÇA — equipe", t["dif"], "ok" if abs(t["dif"]) <= max(LIMIAR_RS, 0.01 * t["devido"]) else "bad"),
                      ("Pago sem base no Trinks (gerente etc.)", t["sem_base"], "warn")):
        c(r, 1, rot, bold=True, fill=f); c(r, 2, v, bold=True, num=True, fill=f); r += 1
    r += 1
    hdr = ["Profissional", "Devido (Trinks)", "Pago (Pix)", "Diferença"] + [f"Devido {q}" for q in qs] + ["Situação", "Pergunta", "Pix pagos", "Resposta do Lucas"]
    for i, h in enumerate(hdr, 1):
        c(r, i, h, bold=True, fill="hdr", color="FFFFFF").alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30
    r0 = r + 1
    for ln in res["linhas"]:
        r += 1
        porq = {x["q"]: x["devido"] for x in ln["quinzenas"]}
        fill = "ok" if ln["situacao"] == "OK" else ("warn" if ln["situacao"] in ("Cadastro excluido no Trinks", "Pix sem producao no Trinks", "Sem base no Trinks") else "bad")
        vals = [ln["nome"], ln["devido"], ln["pago"], ln["dif"]] + [porq.get(q, 0) for q in qs] + \
               [ln["situacao"], ln["pergunta"], " · ".join(ln["pix"]), respostas.get(ln["nome"])]
        for i, v in enumerate(vals, 1):
            c(r, i, v, num=(2 <= i <= 4 + len(qs)), fill=fill if i in (1, 4, 5 + len(qs)) else None, bold=(i == 4))
        ws.cell(r, len(vals)).fill = F["amar"]
        ws.row_dimensions[r].height = max(30, 15 * (max(len(ln["pergunta"]), len(" · ".join(ln["pix"]))) // 48 + 1))
    r += 1
    c(r, 1, "TOTAL", bold=True, fill="tot")
    for col in range(2, 5 + len(qs)):
        L = openpyxl.utils.get_column_letter(col)
        c(r, col, f"=SUM({L}{r0}:{L}{r - 1})", bold=True, fill="tot", num=True)
    ws.freeze_panes = ws.cell(r0, 2)
    wb.save(PEND)


def main():
    seco, forcar = "--seco" in sys.argv, "--forcar" in sys.argv
    try:
        res = conciliar()
    except PermissionError:
        print("Comissoes x Trinks: Excel do painel inacessivel agora; tento na proxima rodada.")
        return 0
    t = res["totais"]
    pend = [l for l in res["linhas"] if l["situacao"] != "OK"]
    resumo = (f"Comissoes x Trinks {res['periodo']}: devido R$ {t['devido']:.2f} x pago R$ {t['pago']:.2f} "
              f"(dif R$ {t['dif']:.2f}) - {len(pend)} pendencia(s)")
    if seco:
        print(ascii_(resumo))
        for l in pend[:15]:
            print(ascii_(f"  {l['situacao'][:28]:28} {l['nome'][:30]:30} devido {l['devido']:>9.2f} pago {l['pago']:>9.2f} dif {l['dif']:>9.2f}"))
        return 0
    assinatura = hashlib.sha1(json.dumps({k: res[k] for k in ("totais", "linhas")}, sort_keys=True).encode()).hexdigest()
    anterior = (ler_json(RESULTADO, {}) or {}).get("assinatura")
    if anterior == assinatura and not forcar:
        print(ascii_(resumo + " - sem mudanca, nao regravei"))
        return 0
    lock = PASTA / ("~$" + PEND.name)
    if lock.exists():
        print(ascii_(resumo + " - arquivo de pendencias aberto, gravo na proxima rodada"))
        return 0
    try:
        gravar_aba(res)
    except PermissionError:
        print(ascii_(resumo + " - arquivo de pendencias em uso, gravo na proxima rodada"))
        return 0
    res["assinatura"] = assinatura
    RESULTADO.write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    print(ascii_(resumo + " - aba atualizada"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
