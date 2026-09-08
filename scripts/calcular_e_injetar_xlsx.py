# -*- coding: utf-8 -*-
"""Avalia as fórmulas do .xlsx, confere contra o modelo em Python e grava os
resultados como valores em cache dentro do próprio arquivo.

Por que existe: o openpyxl escreve fórmulas sem valor calculado. O Excel recalcula
ao abrir, mas visualizadores rápidos (prévia do e-mail, do celular, do Drive) leem
só o cache — e mostram a planilha em branco. Aqui as fórmulas continuam no arquivo,
vivas e editáveis; ganham apenas um valor inicial para quem só quer olhar.

    python calcular_e_injetar.py [arquivo.xlsx]
"""
from __future__ import annotations

import re
import shutil
import sys
import zipfile
import pathlib
from xml.etree import ElementTree as ET

import formulas

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NSR = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSPKG = "http://schemas.openxmlformats.org/package/2006/relationships"
ERRO = re.compile(r"^#(REF!|NAME\?|VALUE!|DIV/0!|N/A|NUM!|NULL!)$")

ARQ = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "FAST-Insights-modelo-financeiro.xlsx")


# ---------------------------------------------------------------- 1. calcular
def calcular(caminho: pathlib.Path) -> dict[tuple[str, str], object]:
    xl = formulas.ExcelModel().loads(str(caminho)).finish()
    sol = xl.calculate()
    fora: dict[tuple[str, str], object] = {}
    for chave, val in sol.items():
        m = re.match(r"^'\[.*?\](.+?)'!([A-Z]+\d+)$", chave)
        if not m:
            continue
        try:
            v = val.value[0, 0]
        except Exception:
            v = getattr(val, "value", val)
        fora[(m.group(1).upper(), m.group(2))] = v
    return fora


# ---------------------------------------------------------------- 2. conferir
ESPERADO_GERAL = [
    ("PREMISSAS", "B12", 319.0, 0.01, "ticket médio"),
    ("PREMISSAS", "B19", 45.0, 0.01, "custo variável"),
    ("PREMISSAS", "B20", 274.0, 0.01, "margem de contribuição"),
    ("MODELO MENSAL", "M40", 953229, 2, "caixa no mês 36"),
    ("RESUMO", "B7", 267088, 2, "capital necessário"),
    ("RESUMO", "B9", 15, 0, "1º mês positivo"),
    ("RESUMO", "B10", 23, 0, "mês em que o caixa zera"),
    ("RESUMO", "B17", 1526137, 2, "receita ano 3"),
    ("RESUMO", "C17", 883248, 2, "resultado ano 3"),
    ("RESUMO", "B20", 127.74, 0.05, "ponto de equilíbrio"),
    ("RESUMO", "B21", 91.24, 0.05, "ponto de equilíbrio com painel"),
    ("RESUMO", "B30", 1.46, 0.02, "payback do CAC"),
    ("RESUMO", "B32", 45.67, 0.05, "LTV / CAC"),
]

# (nome, 1º mês positivo, caixa zera, capital, receita ano 3, resultado ano 3, lojas)
ESPERADO_CEN = [
    ("Base",                       15, 23, 267088, 1526137,  883248, 389),
    ("Adoção metade",              19, 35, 327402,  823069,  291624, 195),
    ("Atraso 6 meses",             21, 29, 292811, 1417651,  752896, 386),
    ("Preço 20% menor",            17, 27, 302405, 1244028,  601139, 389),
    ("Sem Painel da Rede",         17, 26, 335658, 1406137,  763248, 389),
    ("Participação 20%",           17, 28, 318259, 1220910,  578021, 389),
    ("Participação 35%",           19, 35, 377326,  991989,  349100, 389),
    ("Participação 50%",           23,  0, 474338,  763069,  120180, 389),
    ("Pior caso combinado",        33,  0, 500221,  499736,  -42641, 193),
    ("Enxuta",                     13, 21, 214871, 1526137, 1003248, 389),
    ("OBRIGATÓRIO sem particip.",  10, 16, 178701, 1748830, 1051296, 570),
    ("OBRIGATÓRIO particip. 20%",  12, 19, 200150, 1399064,  701530, 570),
    ("OBRIGATÓRIO particip. 30%",  13, 21, 215809, 1224181,  526647, 570),
    ("OBRIGATÓRIO 299 · 30%",      12, 18, 198024, 1453133,  755599, 570),
]
LIN_RES = 21


def conferir(vals) -> bool:
    erros = [(a, b, v) for (a, b), v in vals.items() if ERRO.match(str(v).strip())]
    print(f"CÉLULAS COM ERRO: {len(erros)}")
    for a, b, v in erros[:10]:
        print("   ", a, b, "=", v)

    def num(aba, cel):
        try:
            return float(vals.get((aba, cel), "nan"))
        except Exception:
            return float("nan")

    ok_total = not erros
    print("\nGERAL")
    for aba, cel, esp, tol, rot in ESPERADO_GERAL:
        got = num(aba, cel)
        ok = abs(got - esp) <= tol
        ok_total &= ok
        print(f"  {'OK ' if ok else 'ERRO'} {rot:32} {got:>16,.2f}   esperado {esp:>14,.2f}"
              .replace(",", "."))

    print("\nCENÁRIOS")
    for i, (nome, *esp) in enumerate(ESPERADO_CEN):
        r = LIN_RES + i
        got = [num("CENÁRIOS", f"{c}{r}") for c in "BCDEFG"]
        ok = all(abs(g - e) <= max(2, abs(e) * 0.001) for g, e in zip(got, esp))
        ok_total &= ok
        f = lambda x: f"{x:,.0f}".replace(",", ".")
        print(f"  {'OK ' if ok else 'ERRO'} {nome:26} 1º+ {f(got[0]):>3} | cx0 {f(got[1]):>3} | "
              f"cap {f(got[2]):>9} | rec3 {f(got[3]):>10} | res3 {f(got[4]):>9} | "
              f"lojas {f(got[5]):>4}")
    return ok_total


# ---------------------------------------------------------------- 3. injetar
def mapa_planilhas(z: zipfile.ZipFile) -> dict[str, str]:
    """nome da aba (maiúsculas) -> caminho do XML dentro do zip."""
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    alvo = {r.get("Id"): r.get("Target") for r in rels.findall(f"{{{NSPKG}}}Relationship")}
    fora = {}
    for sh in wb.find(f"{{{NS}}}sheets"):
        destino = alvo[sh.get(f"{{{NSR}}}id")].lstrip("/")
        if not destino.startswith("xl/"):
            destino = "xl/" + destino
        fora[sh.get("name").upper()] = destino
    return fora


def injetar(caminho: pathlib.Path, vals) -> int:
    ET.register_namespace("", NS)
    ET.register_namespace("r", NSR)
    tmp = caminho.with_suffix(".tmp.xlsx")
    gravadas = 0
    with zipfile.ZipFile(caminho) as z:
        planilhas = mapa_planilhas(z)
        xml_de = {v: k for k, v in planilhas.items()}
        itens = {n: z.read(n) for n in z.namelist()}

    for nome_xml, aba in xml_de.items():
        raiz = ET.fromstring(itens[nome_xml])
        mudou = False
        for c in raiz.iter(f"{{{NS}}}c"):
            f = c.find(f"{{{NS}}}f")
            if f is None:
                continue
            v = vals.get((aba, c.get("r")))
            if v is None:
                continue
            for antigo in c.findall(f"{{{NS}}}v"):
                c.remove(antigo)
            novo = ET.SubElement(c, f"{{{NS}}}v")
            texto = str(v)
            if isinstance(v, bool):
                c.set("t", "b")
                novo.text = "1" if v else "0"
            elif ERRO.match(texto.strip()):
                c.set("t", "e")
                novo.text = texto.strip()
            else:
                try:
                    novo.text = repr(float(v))
                    c.attrib.pop("t", None)
                except (TypeError, ValueError):
                    if texto == "":
                        c.remove(novo)
                        continue
                    c.set("t", "str")
                    novo.text = texto
            gravadas += 1
            mudou = True
        if mudou:
            itens[nome_xml] = ET.tostring(raiz, encoding="UTF-8", xml_declaration=True)

    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, dados in itens.items():
            z.writestr(nome, dados)
    shutil.move(tmp, caminho)
    return gravadas


def main() -> int:
    print(f"Arquivo: {ARQ}\n")
    vals = calcular(ARQ)
    ok = conferir(vals)
    n = injetar(ARQ, vals)
    print(f"\nValores em cache gravados: {n}")

    from openpyxl import load_workbook
    wbv = load_workbook(ARQ, data_only=True)
    wbf = load_workbook(ARQ)
    vazias = []
    for aba in wbv.sheetnames:
        preenchidas = sum(1 for row in wbv[aba].iter_rows()
                          for c in row if c.value not in (None, ""))
        vazias.append((aba, preenchidas))
    print("\nCÉLULAS COM VALOR VISÍVEL POR ABA (o que um visualizador mostra)")
    for aba, n_ in vazias:
        print(f"  {'OK ' if n_ else 'VAZIA'} {aba:16} {n_:>6}")
    assert wbf["Modelo mensal"]["M40"].value.startswith("="), "as fórmulas sumiram!"
    print("\nFórmulas preservadas: Modelo mensal!M40 =", wbf["Modelo mensal"]["M40"].value)
    print("\nRESULTADO:", "tudo confere" if ok and all(n_ for _, n_ in vazias) else "REVISAR")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
